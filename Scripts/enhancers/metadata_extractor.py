from typing import Dict, Any, Optional, List
import os
from datetime import datetime
import hashlib
import magic
import PyPDF2
from docx import Document
from pptx import Presentation
import exifread
from pathlib import Path
import logging
from dataclasses import dataclass
import json
import mimetypes
from langdetect import detect
import spacy
from collections import Counter
import textstat

logger = logging.getLogger(__name__)

@dataclass
class MetadataConfig:
    extract_content_stats: bool = True
    extract_file_stats: bool = True
    extract_language_info: bool = True
    extract_readability_metrics: bool = True
    compute_checksums: bool = True
    extract_named_entities: bool = True

class MetadataExtractor:
    def __init__(self, config: MetadataConfig = None):
        self.config = config or MetadataConfig()
        if self.config.extract_named_entities:
            try:
                self.nlp = spacy.load("en_core_web_sm")
            except OSError:
                logger.warning("Spacy model not found. Installing en_core_web_sm...")
                import subprocess
                subprocess.run(["python", "-m", "spacy", "download", "en_core_web_sm"])
                self.nlp = spacy.load("en_core_web_sm")

    def extract_metadata(self, file_path: str) -> Dict[str, Any]:
        """Extract comprehensive metadata from a file"""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        metadata = {
            "file_path": file_path,
            "filename": os.path.basename(file_path),
            "extraction_time": datetime.now().isoformat()
        }

        # Basic file metadata
        if self.config.extract_file_stats:
            metadata.update(self._get_file_stats(file_path))

        # File type specific metadata
        file_type_metadata = self._get_file_type_metadata(file_path)
        if file_type_metadata:
            metadata["format_metadata"] = file_type_metadata

        # Content statistics if text content is available
        if self.config.extract_content_stats:
            content = self._extract_text_content(file_path)
            if content:
                metadata.update(self._analyze_content(content))

        return metadata

    def _get_file_stats(self, file_path: str) -> Dict[str, Any]:
        """Get basic file statistics"""
        stats = os.stat(file_path)
        mime_type = magic.from_file(file_path, mime=True)
        
        file_stats = {
            "size_bytes": stats.st_size,
            "created_time": datetime.fromtimestamp(stats.st_ctime).isoformat(),
            "modified_time": datetime.fromtimestamp(stats.st_mtime).isoformat(),
            "accessed_time": datetime.fromtimestamp(stats.st_atime).isoformat(),
            "mime_type": mime_type,
            "extension": os.path.splitext(file_path)[1].lower()
        }

        if self.config.compute_checksums:
            file_stats.update(self._compute_checksums(file_path))

        return file_stats

    def _compute_checksums(self, file_path: str) -> Dict[str, str]:
        """Compute multiple checksums for a file"""
        checksums = {}
        algorithms = {
            'md5': hashlib.md5(),
            'sha1': hashlib.sha1(),
            'sha256': hashlib.sha256()
        }

        with open(file_path, 'rb') as f:
            while chunk := f.read(8192):
                for hash_obj in algorithms.values():
                    hash_obj.update(chunk)

        return {
            f"{algo}_hash": hash_obj.hexdigest()
            for algo, hash_obj in algorithms.items()
        }

    def _get_file_type_metadata(self, file_path: str) -> Optional[Dict[str, Any]]:
        """Extract metadata specific to file type"""
        ext = os.path.splitext(file_path)[1].lower()
        
        try:
            if ext == '.pdf':
                return self._extract_pdf_metadata(file_path)
            elif ext in ['.docx', '.doc']:
                return self._extract_docx_metadata(file_path)
            elif ext in ['.pptx', '.ppt']:
                return self._extract_pptx_metadata(file_path)
            elif ext in ['.jpg', '.jpeg', '.png', '.gif']:
                return self._extract_image_metadata(file_path)
            else:
                return None
        except Exception as e:
            logger.warning(f"Failed to extract format-specific metadata: {e}")
            return None

    def _extract_pdf_metadata(self, file_path: str) -> Dict[str, Any]:
        """Extract metadata from PDF files"""
        with open(file_path, 'rb') as f:
            pdf = PyPDF2.PdfReader(f)
            info = pdf.metadata
            return {
                "page_count": len(pdf.pages),
                "title": info.get('/Title', ''),
                "author": info.get('/Author', ''),
                "subject": info.get('/Subject', ''),
                "creator": info.get('/Creator', ''),
                "producer": info.get('/Producer', ''),
                "creation_date": info.get('/CreationDate', ''),
                "modification_date": info.get('/ModDate', '')
            }

    def _extract_docx_metadata(self, file_path: str) -> Dict[str, Any]:
        """Extract metadata from DOCX files"""
        doc = Document(file_path)
        core_props = doc.core_properties
        return {
            "author": core_props.author or '',
            "title": core_props.title or '',
            "created": core_props.created.isoformat() if core_props.created else '',
            "modified": core_props.modified.isoformat() if core_props.modified else '',
            "last_modified_by": core_props.last_modified_by or '',
            "revision": core_props.revision or 0,
            "category": core_props.category or '',
            "comments": core_props.comments or '',
            "paragraph_count": len(doc.paragraphs),
            "word_count": sum(len(p.text.split()) for p in doc.paragraphs)
        }

    def _extract_pptx_metadata(self, file_path: str) -> Dict[str, Any]:
        """Extract metadata from PPTX files"""
        prs = Presentation(file_path)
        return {
            "slide_count": len(prs.slides),
            "slide_heights": prs.slide_height,
            "slide_width": prs.slide_width
        }

    def _extract_image_metadata(self, file_path: str) -> Dict[str, Any]:
        """Extract metadata from image files"""
        with open(file_path, 'rb') as f:
            tags = exifread.process_file(f)
            return {str(k): str(v) for k, v in tags.items()}

    def _extract_text_content(self, file_path: str) -> Optional[str]:
        """Extract text content from supported file types"""
        ext = os.path.splitext(file_path)[1].lower()
        
        try:
            if ext == '.pdf':
                with open(file_path, 'rb') as f:
                    pdf = PyPDF2.PdfReader(f)
                    return ' '.join(page.extract_text() for page in pdf.pages)
            elif ext == '.docx':
                doc = Document(file_path)
                return ' '.join(paragraph.text for paragraph in doc.paragraphs)
            elif ext in ['.txt', '.md', '.json', '.xml', '.yml', '.yaml']:
                with open(file_path, 'r', encoding='utf-8') as f:
                    return f.read()
            else:
                return None
        except Exception as e:
            logger.warning(f"Failed to extract text content: {e}")
            return None

    def _analyze_content(self, content: str) -> Dict[str, Any]:
        """Analyze text content for various metrics"""
        analysis = {}

        if self.config.extract_language_info:
            try:
                analysis["language"] = detect(content)
            except:
                analysis["language"] = "unknown"

        if self.config.extract_readability_metrics:
            analysis["readability"] = {
                "flesch_reading_ease": textstat.flesch_reading_ease(content),
                "flesch_kincaid_grade": textstat.flesch_kincaid_grade(content),
                "gunning_fog": textstat.gunning_fog(content),
                "smog_index": textstat.smog_index(content),
                "automated_readability_index": textstat.automated_readability_index(content),
                "coleman_liau_index": textstat.coleman_liau_index(content),
                "dale_chall_readability_score": textstat.dale_chall_readability_score(content)
            }

        if self.config.extract_named_entities and self.nlp:
            doc = self.nlp(content[:100000])  # Limit for performance
            entities = Counter((ent.text.strip(), ent.label_) for ent in doc.ents)
            analysis["named_entities"] = [
                {"text": text, "label": label, "count": count}
                for (text, label), count in entities.most_common(50)
            ]

        # Basic text statistics
        words = content.split()
        sentences = content.split('.')
        analysis.update({
            "word_count": len(words),
            "character_count": len(content),
            "sentence_count": len(sentences),
            "average_word_length": sum(len(word) for word in words) / len(words) if words else 0,
            "average_sentence_length": len(words) / len(sentences) if sentences else 0
        })

        return analysis

    def batch_extract_metadata(self, file_paths: List[str]) -> Dict[str, Dict[str, Any]]:
        """Extract metadata from multiple files in parallel"""
        from concurrent.futures import ThreadPoolExecutor
        
        results = {}
        with ThreadPoolExecutor() as executor:
            future_to_path = {
                executor.submit(self.extract_metadata, path): path 
                for path in file_paths
            }
            for future in concurrent.futures.as_completed(future_to_path):
                path = future_to_path[future]
                try:
                    metadata = future.result()
                    results[path] = metadata
                except Exception as e:
                    logger.error(f"Error processing {path}: {e}")
                    results[path] = {"error": str(e)}
        
        return results
